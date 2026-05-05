"""Layout-related tools for paragraph, page setup, and header/footer operations."""

from typing import Optional
from docx import Document
from docx.shared import Pt

from word_document_server.utils.file_utils import ensure_docx_extension


def _pt_to_float(value):
    if value is None:
        return None
    return round(value.pt, 2)


async def set_paragraph_layout(
    filename: str,
    paragraph_index: int,
    line_spacing: Optional[float] = None,
    first_line_indent_pt: Optional[float] = None,
) -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if len(doc.paragraphs) == 0 and paragraph_index == 0:
        doc.add_paragraph("")

    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return f"Paragraph index {paragraph_index} out of range. Document has {len(doc.paragraphs)} paragraphs."

    paragraph = doc.paragraphs[paragraph_index]
    pf = paragraph.paragraph_format

    if line_spacing is not None:
        pf.line_spacing = float(line_spacing)

    if first_line_indent_pt is not None:
        pf.first_line_indent = Pt(float(first_line_indent_pt))

    doc.save(filename)
    return f"Updated paragraph {paragraph_index} layout in {filename}"


async def get_paragraph_layout(filename: str, paragraph_index: int) -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if len(doc.paragraphs) == 0 and paragraph_index == 0:
        doc.add_paragraph("")

    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return f"Paragraph index {paragraph_index} out of range. Document has {len(doc.paragraphs)} paragraphs."

    paragraph = doc.paragraphs[paragraph_index]
    pf = paragraph.paragraph_format

    info = {
        "paragraph_index": paragraph_index,
        "text_preview": paragraph.text[:80],
        "line_spacing": pf.line_spacing,
        "first_line_indent_pt": _pt_to_float(pf.first_line_indent),
        "left_indent_pt": _pt_to_float(pf.left_indent),
        "right_indent_pt": _pt_to_float(pf.right_indent),
        "space_before_pt": _pt_to_float(pf.space_before),
        "space_after_pt": _pt_to_float(pf.space_after),
    }
    return str(info)


async def set_page_setup(
    filename: str,
    section_index: int = 0,
    top_margin_pt: Optional[float] = None,
    bottom_margin_pt: Optional[float] = None,
    left_margin_pt: Optional[float] = None,
    right_margin_pt: Optional[float] = None,
    page_width_pt: Optional[float] = None,
    page_height_pt: Optional[float] = None,
) -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if section_index < 0 or section_index >= len(doc.sections):
        return f"Section index {section_index} out of range. Document has {len(doc.sections)} sections."

    sec = doc.sections[section_index]
    if top_margin_pt is not None:
        sec.top_margin = Pt(float(top_margin_pt))
    if bottom_margin_pt is not None:
        sec.bottom_margin = Pt(float(bottom_margin_pt))
    if left_margin_pt is not None:
        sec.left_margin = Pt(float(left_margin_pt))
    if right_margin_pt is not None:
        sec.right_margin = Pt(float(right_margin_pt))
    if page_width_pt is not None:
        sec.page_width = Pt(float(page_width_pt))
    if page_height_pt is not None:
        sec.page_height = Pt(float(page_height_pt))

    doc.save(filename)
    return f"Updated section {section_index} page setup in {filename}"


async def get_page_setup(filename: str, section_index: int = 0) -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if section_index < 0 or section_index >= len(doc.sections):
        return f"Section index {section_index} out of range. Document has {len(doc.sections)} sections."

    sec = doc.sections[section_index]
    info = {
        "section_index": section_index,
        "top_margin_pt": _pt_to_float(sec.top_margin),
        "bottom_margin_pt": _pt_to_float(sec.bottom_margin),
        "left_margin_pt": _pt_to_float(sec.left_margin),
        "right_margin_pt": _pt_to_float(sec.right_margin),
        "page_width_pt": _pt_to_float(sec.page_width),
        "page_height_pt": _pt_to_float(sec.page_height),
    }
    return str(info)


async def set_header_footer(
    filename: str,
    text: str,
    section_index: int = 0,
    target: str = "header",
    clear_existing: bool = True,
) -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if section_index < 0 or section_index >= len(doc.sections):
        return f"Section index {section_index} out of range. Document has {len(doc.sections)} sections."

    section = doc.sections[section_index]
    container = section.header if target.lower() == "header" else section.footer

    if clear_existing:
        for p in list(container.paragraphs):
            p._element.getparent().remove(p._element)

    if container.paragraphs:
        container.paragraphs[0].text = text
    else:
        container.add_paragraph(text)

    doc.save(filename)
    return f"Updated {target.lower()} text for section {section_index} in {filename}"


async def get_header_footer(filename: str, section_index: int = 0, target: str = "header") -> str:
    filename = ensure_docx_extension(filename)
    doc = Document(filename)

    if section_index < 0 or section_index >= len(doc.sections):
        return f"Section index {section_index} out of range. Document has {len(doc.sections)} sections."

    section = doc.sections[section_index]
    container = section.header if target.lower() == "header" else section.footer
    texts = [p.text for p in container.paragraphs if p.text.strip()]
    return str({"section_index": section_index, "target": target.lower(), "text": "\n".join(texts)})
